import argparse
import csv
import io
import json
import re
import statistics
from collections import defaultdict
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Inches, Pt  # noqa: E402

DB_MARKERS = ("db", "postgres")


def num(value) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_text(path: Path) -> str:
    """Locust на Windows записує CSV у системному кодуванні (cp1251), тому пробуємо кілька."""
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def read_csv(path: Path) -> list[dict]:
    if not path.exists():
        return []
    return list(csv.DictReader(io.StringIO(read_text(path), newline="")))


def ms(value: float | None) -> str:
    return "—" if value is None else f"{value:,.0f}".replace(",", " ")


def f1(value: float | None) -> str:
    return "—" if value is None else f"{value:,.1f}".replace(",", " ")


class Report:
    """Один опис звіту → два формати (Markdown і Word)."""

    def __init__(self) -> None:
        self.blocks: list[tuple] = []

    def h(self, level: int, text: str) -> None:
        self.blocks.append(("h", level, text))

    def p(self, text: str) -> None:
        self.blocks.append(("p", text))

    def bullets(self, items: list[str]) -> None:
        self.blocks.append(("ul", items))

    def table(self, headers: list[str], rows: list[list[str]]) -> None:
        self.blocks.append(("table", headers, rows))

    def image(self, path: Path, caption: str) -> None:
        self.blocks.append(("img", path, caption))

    def to_md(self) -> str:
        out: list[str] = []
        for block in self.blocks:
            kind = block[0]
            if kind == "h":
                out.append(f"{'#' * block[1]} {block[2]}\n")
            elif kind == "p":
                out.append(block[1] + "\n")
            elif kind == "ul":
                out.append("\n".join(f"- {i}" for i in block[1]) + "\n")
            elif kind == "table":
                out.append("| " + " | ".join(block[1]) + " |")
                out.append("|" + "---|" * len(block[1]))
                out.extend("| " + " | ".join(r) + " |" for r in block[2])
                out.append("")
            elif kind == "img":
                out.append(f"![{block[2]}]({block[1].name})\n\n*{block[2]}*\n")
        return "\n".join(out)

    @staticmethod
    def _runs(paragraph, text: str) -> None:
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            if part:
                paragraph.add_run(part).bold = i % 2 == 1

    def to_docx(self, path: Path) -> None:
        doc = Document()
        zoom = doc.settings.element.find(qn("w:zoom"))
        if zoom is not None:  # у стандартному шаблоні python-docx бракує обов'язкового атрибута
            zoom.set(qn("w:percent"), "100")
        style = doc.styles["Normal"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        for block in self.blocks:
            kind = block[0]
            if kind == "h":
                doc.add_heading(block[2], level=block[1] - 1 if block[1] > 1 else 0)
            elif kind == "p":
                self._runs(doc.add_paragraph(), block[1])
            elif kind == "ul":
                for item in block[1]:
                    self._runs(doc.add_paragraph(style="List Bullet"), item)
            elif kind == "table":
                table = doc.add_table(rows=1, cols=len(block[1]), style="Table Grid")
                for cell, title in zip(table.rows[0].cells, block[1], strict=False):
                    cell.text = ""
                    cell.paragraphs[0].add_run(title).bold = True
                for row in block[2]:
                    for cell, value in zip(table.add_row().cells, row, strict=False):
                        cell.text = value
                if len(block[1]) >= 5:  # широкі таблиці: перша колонка (назви) ширша
                    table.autofit = False
                    first = Inches(2.0)
                    rest = Inches((6.5 - 2.0) / (len(block[1]) - 1))
                    for i, column in enumerate(table.columns):
                        column.width = first if i == 0 else rest  # сітка таблиці (для LibreOffice)
                    for row in table.rows:
                        for i, cell in enumerate(row.cells):
                            cell.width = first if i == 0 else rest
                doc.add_paragraph()
            elif kind == "img":
                doc.add_picture(str(block[1]), width=Inches(6.2))
                caption = doc.add_paragraph()
                caption.add_run(block[2]).italic = True
        doc.save(str(path))


# ---------------------------------------------------------------- графіки


def draw_charts(history: list[dict], resources: list[dict], out: Path) -> dict[str, Path]:
    charts: dict[str, Path] = {}
    if not history:
        return charts
    t0 = history[0]["t"]
    xs = [r["t"] - t0 for r in history]
    users = [r["users"] for r in history]

    def with_users(ax):
        ax2 = ax.twinx()
        ax2.fill_between(xs, users, step="post", color="#999999", alpha=0.15)
        ax2.set_ylabel("Віртуальні користувачі (VU)")
        ax2.set_ylim(bottom=0)
        return ax2

    fig, ax = plt.subplots(figsize=(9, 4.2))
    ax.plot(xs, [r["p50"] for r in history], label="Медіана", color="#2b50e0")
    ax.plot(xs, [r["p95"] for r in history], label="95-й перцентиль", color="#b3261e")
    ax.set_xlabel("Час від початку тесту, с")
    ax.set_ylabel("Час відповіді, мс")
    ax.set_title("Динаміка часу відповіді (вікно 10 с)")
    ax.grid(alpha=0.3)
    ax.legend(loc="upper left")
    with_users(ax)
    fig.tight_layout()
    charts["response"] = out / "chart_response_time.png"
    fig.savefig(charts["response"], dpi=130)
    plt.close(fig)

    fig, ax = plt.subplots(figsize=(9, 4.2))
    ax.plot(xs, [r["rps"] for r in history], label="Запитів/с (RPS)", color="#1f6b4a")
    ax.plot(xs, [r["fails"] for r in history], label="Помилок/с", color="#b3261e")
    ax.set_xlabel("Час від початку тесту, с")
    ax.set_ylabel("Запитів за секунду")
    ax.set_title("Пропускна здатність")
    ax.grid(alpha=0.3)
    ax.legend(loc="upper left")
    with_users(ax)
    fig.tight_layout()
    charts["throughput"] = out / "chart_throughput.png"
    fig.savefig(charts["throughput"], dpi=130)
    plt.close(fig)

    if resources:
        by_name: dict[str, list[dict]] = defaultdict(list)
        for row in resources:
            by_name[row["name"]].append(row)
        fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(9, 5.6), sharex=True)
        for name, rows in by_name.items():
            rx = [num(r["timestamp"]) - t0 for r in rows]
            ax1.plot(rx, [num(r["cpu_percent"]) for r in rows], label=name)
            ax2.plot(rx, [num(r["mem_mib"]) for r in rows], label=name)
        ax1.set_ylabel("CPU, % від одного ядра")
        ax2.set_ylabel("RAM, МіБ")
        ax2.set_xlabel("Час від початку тесту, с")
        ax1.set_title("Ресурси сервера під час тесту")
        for ax in (ax1, ax2):
            ax.grid(alpha=0.3)
            ax.legend(loc="upper left")
            ax.set_xlim(0, xs[-1])
        fig.tight_layout()
        charts["resources"] = out / "chart_resources.png"
        fig.savefig(charts["resources"], dpi=130)
        plt.close(fig)
    return charts


# ---------------------------------------------------------------- аналіз


def stage_table(history: list[dict]) -> list[dict]:
    """Підсумок по рівнях навантаження: медіани RPS і затримок для кожної кількості VU."""
    groups: dict[int, list[dict]] = defaultdict(list)
    for r in history:
        if r["p95"] is not None and r["p50"] is not None:
            groups[round(r["users"] / 5) * 5].append(r)
    return [
        {
            "users": users,
            "samples": len(items),
            "rps": statistics.median(i["rps"] for i in items),
            "p50": statistics.median(i["p50"] for i in items),
            "p95": statistics.median(i["p95"] for i in items),
            "fails": statistics.median(i["fails"] for i in items),
        }
        for users, items in sorted(groups.items())
        if users > 0 and len(items) >= 5
    ]


def find_knee(history: list[dict]) -> dict | None:
    """Перший рівень, на якому медіанний p95 перевищує 3× базовий (і 500 мс)."""
    stages = stage_table(history)
    if len(stages) < 2:
        return None
    limit = max(stages[0]["p95"] * 3, 500)
    for prev, cur in zip(stages, stages[1:], strict=False):
        if cur["p95"] > limit:
            return {"baseline": stages[0]["p95"], "previous": prev, "stage": cur}
    return None


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--results", default="loadtests/results")
    args = parser.parse_args()
    res = Path(args.results)

    meta = json.loads((res / "meta.json").read_text(encoding="utf-8"))
    stats = read_csv(res / "run_stats.csv")
    failures = read_csv(res / "run_failures.csv")
    resources = read_csv(res / "resources.csv")
    server_log = {}
    if (res / "server_log.json").exists():
        server_log = json.loads((res / "server_log.json").read_text(encoding="utf-8"))
    double = None
    if (res / "double_booking.json").exists():
        double = json.loads((res / "double_booking.json").read_text(encoding="utf-8"))

    history = []
    for row in read_csv(res / "run_stats_history.csv"):
        if row["Name"] != "Aggregated" or num(row["Timestamp"]) is None:
            continue
        history.append(
            {
                "t": num(row["Timestamp"]),
                "users": int(num(row["User Count"]) or 0),
                "rps": num(row["Requests/s"]) or 0.0,
                "fails": num(row["Failures/s"]) or 0.0,
                "p50": num(row["50%"]),
                "p95": num(row["95%"]),
            }
        )
    history = [r for r in history if r["users"] > 0]
    if not stats or not history:
        raise SystemExit("Немає даних Locust у папці results — спершу виконайте тест.")

    total = next(r for r in stats if r["Name"] == "Aggregated")
    endpoints = sorted((r for r in stats if r["Name"] != "Aggregated"), key=lambda r: r["Name"])
    requests_total = int(total["Request Count"])
    failures_total = int(total["Failure Count"])
    duration = history[-1]["t"] - history[0]["t"]
    error_rate = failures_total / requests_total * 100 if requests_total else 0.0
    peak_rps = max(r["rps"] for r in history)
    max_users = max(r["users"] for r in history)

    charts = draw_charts(history, resources, res)

    # ресурси
    res_rows = []
    by_name: dict[str, list[dict]] = defaultdict(list)
    for row in resources:
        by_name[row["name"]].append(row)
    app_cpu_max = db_cpu_max = 0.0
    for name, rows in sorted(by_name.items()):
        cpu = [num(r["cpu_percent"]) for r in rows]
        mem = [num(r["mem_mib"]) for r in rows]
        res_rows.append(
            [name, f1(statistics.mean(cpu)), f1(max(cpu)), f1(statistics.mean(mem)), f1(max(mem))]
        )
        if any(m in name.lower() for m in DB_MARKERS):
            db_cpu_max = max(db_cpu_max, max(cpu))
        elif "redis" not in name.lower():
            app_cpu_max = max(app_cpu_max, max(cpu))

    m = meta
    rep = Report()
    rep.h(1, "Звіт про первинне навантажувальне тестування")
    rep.p(
        f"Проєкт: платформа бронювання коворкінгів та конференц-залів. Дата тесту: {m['started']}. "
        f"Профіль: **{m['profile_title']}**."
    )

    rep.h(2, "1. Мета та вибір інструмента")
    rep.p(
        "Мета — виміряти продуктивність MVP під навантаженням, зафіксувати базові метрики "
        "(RPS, затримки, відсоток помилок, CPU/RAM) і знайти перші вузькі місця."
    )
    rep.p(
        "Інструмент — **Locust**. Причини: сценарій — звичайний Python-код (той самий стек, що й "
        "проєкт), авторизацію й залежності між запитами легко описати програмно; є headless-режим і "
        "CSV-вихід для автоматичного звіту; профіль користувачів задається кодом. "
        "Apache JMeter важчий (GUI, XML-плани), k6 потребує JavaScript."
    )

    rep.h(2, "2. Тестове середовище")
    rep.table(
        ["Параметр", "Значення"],
        [
            ["Середовище застосунку", f"{m['app_environment']} (за відповіддю /health)"],
            ["Адреса", m["host"]],
            ["Застосунок", "FastAPI, один процес uvicorn у Docker-контейнері"],
            ["База даних", "PostgreSQL 16 (окремий контейнер, sandbox-база)"],
            [
                "Комп'ютер",
                f"{m['machine']['os']}, ядер CPU: {m['machine']['cpus']}, RAM: {m['machine']['ram_gib']} ГіБ",
            ],
            ["Інструмент", f"Locust {m['locust_version']}, Python {m['machine']['python']}"],
            ["Збір CPU/RAM", m["monitor_mode"]],
        ],
    )
    rep.p(
        "Генератор навантаження працював на тому ж комп'ютері, що й сервер, тому він теж споживає "
        "частину ресурсів — абсолютні значення слід розглядати як орієнтовні."
    )

    rep.h(2, "3. Цільові ендпоінти та сценарій")
    rep.p(
        "Сценарій імітує типового відвідувача: реєстрація і вхід при старті сесії, далі — перегляд залів, "
        "перевірка зайнятості, бронювання, перегляд і скасування своїх бронювань, з паузами 1–3 с між діями. "
        "90 % віртуальних користувачів поводяться типово (клас CoworkingUser), 10 % — «пікова» поведінка "
        "(PeakUser): вони б'ються за ті самі слоти, що моделює початок тижня."
    )
    rep.table(
        ["Ендпоінт", "Вага в сценарії", "Призначення"],
        [
            ["POST /auth/register", "1 раз на користувача", "реєстрація (хешування пароля)"],
            ["POST /auth/login", "1 раз на користувача", "вхід, отримання JWT"],
            ["GET /", "1", "завантаження сторінки інтерфейсу"],
            ["GET /rooms/", "6", "список залів"],
            ["GET /rooms/{id}/availability", "5", "зайнятість залу за добу"],
            ["GET /bookings/my", "3", "мої бронювання"],
            [
                "POST /bookings/",
                "2 (PeakUser: 8)",
                "створення бронювання, перевірка перетину слотів",
            ],
            ["DELETE /bookings/{id}", "1", "скасування бронювання"],
        ],
    )
    rep.p(
        "Відповідь 409 на бронювання (слот зайнятий) вважається нормальним бізнес-результатом, а не помилкою системи."
    )

    rep.h(2, "4. Профіль навантаження")
    rep.table(
        ["Параметр", "Значення"],
        [
            ["Профіль", m["profile_title"]],
            ["Максимум віртуальних користувачів (VU)", str(max_users)],
            ["Розгін (ramp-up)", m["ramp_text"]],
            ["Тривалість тесту", f"{duration:.0f} с"],
        ],
    )

    rep.h(2, "5. Результати")
    rep.h(3, "5.1. Підсумкові метрики")
    rep.table(
        ["Метрика", "Значення"],
        [
            ["Всього запитів", f"{requests_total}"],
            ["Середній RPS", f1(requests_total / duration if duration else None)],
            ["Пікова пропускна здатність, RPS", f1(peak_rps)],
            ["Час відповіді, середній, мс", ms(num(total["Average Response Time"]))],
            ["Час відповіді, мінімальний, мс", ms(num(total["Min Response Time"]))],
            ["Час відповіді, максимальний, мс", ms(num(total["Max Response Time"]))],
            ["Час відповіді, медіана, мс", ms(num(total["50%"]))],
            ["Час відповіді, 95-й перцентиль, мс", ms(num(total["95%"]))],
            ["Час відповіді, 99-й перцентиль, мс", ms(num(total["99%"]))],
            ["Помилкових відповідей", f"{failures_total}"],
            ["Error rate", f"{error_rate:.2f} %"],
        ],
    )
    stages = stage_table(history)
    if len(stages) >= 2:
        rep.h(3, "5.2. Залежність від кількості віртуальних користувачів")
        rep.table(
            ["VU", "Вимірів", "RPS (медіана)", "p50, мс", "p95, мс", "Помилок/с"],
            [
                [
                    str(st["users"]),
                    str(st["samples"]),
                    f1(st["rps"]),
                    ms(st["p50"]),
                    ms(st["p95"]),
                    f"{st['fails']:.2f}",
                ]
                for st in stages
            ],
        )
        rep.p(
            "Значення — медіани вимірів (вікно 10 с) на кожному рівні навантаження; перехідні моменти розгону відкинуто."
        )
    rep.h(3, "5.3. Метрики по ендпоінтах")
    rep.table(
        ["Ендпоінт", "Запити", "RPS", "Сер., мс", "Мін., мс", "Макс., мс", "p95, мс", "Збоїв"],
        [
            [
                r["Name"],
                r["Request Count"],
                f1(num(r["Requests/s"])),
                ms(num(r["Average Response Time"])),
                ms(num(r["Min Response Time"])),
                ms(num(r["Max Response Time"])),
                ms(num(r["95%"])),
                r["Failure Count"],
            ]
            for r in endpoints
        ],
    )
    rep.h(3, "5.4. Ресурси сервера")
    if res_rows:
        rep.table(
            [
                "Контейнер / процес",
                "CPU сер., %",
                "CPU макс., %",
                "RAM сер., МіБ",
                "RAM макс., МіБ",
            ],
            res_rows,
        )
        rep.p(
            "CPU вимірюється у відсотках від одного ядра (200 % означає два повністю зайняті ядра)."
        )
    else:
        rep.p("Дані про CPU/RAM не збиралися.")

    rep.h(2, "6. Графіки")
    if "response" in charts:
        rep.image(
            charts["response"],
            "Рисунок 1 — Динаміка часу відповіді (медіана та 95-й перцентиль) і кількість VU",
        )
    if "throughput" in charts:
        rep.image(
            charts["throughput"],
            "Рисунок 2 — Пропускна здатність (RPS) і кількість помилок за секунду",
        )
    if "resources" in charts:
        rep.image(charts["resources"], "Рисунок 3 — Завантаження CPU та RAM")

    # ---- вузькі місця
    findings: list[str] = []
    advice: list[str] = []
    knee = find_knee(history)
    if knee:
        cur, prev = knee["stage"], knee["previous"]
        findings.append(
            f"**Межа продуктивності.** При {cur['users']} віртуальних користувачах медіанний 95-й перцентиль "
            f"стрибнув з {ms(prev['p95'])} мс (на {prev['users']} VU) до {ms(cur['p95'])} мс, а пропускна здатність "
            f"змінилася з {prev['rps']:.1f} до {cur['rps']:.1f} RPS — система перестала справлятися з потоком запитів."
        )
    elif len(stages) < 2:
        findings.append(
            "**Профіль має один рівень навантаження**, тому межу продуктивності не визначено. "
            "Для її пошуку виконайте профіль stress (ступінчасте навантаження до 200 VU)."
        )
    else:
        findings.append(
            f"**Різкого зростання затримки не зафіксовано** до {max_users} VU: запас міцності не вичерпано. "
            "Збільште навантаження, щоб знайти межу."
        )

    ranked = sorted(
        (r for r in endpoints if num(r["95%"]) is not None), key=lambda r: -num(r["95%"])
    )
    if ranked:
        top = ", ".join(f"{r['Name']} (p95 = {ms(num(r['95%']))} мс)" for r in ranked[:3])
        findings.append(f"**Найповільніші ендпоінти за p95:** {top}.")

    by_ep = {r["Name"]: r for r in endpoints}
    light = by_ep.get("GET /rooms/")
    heavy = [by_ep[n] for n in ("POST /auth/login", "POST /auth/register") if n in by_ep]
    if light and heavy and max(num(h["95%"]) or 0 for h in heavy) > 3 * (num(light["95%"]) or 1):
        slow = max(heavy, key=lambda h: num(h["95%"]) or 0)
        findings.append(
            f"**Автентифікація набагато повільніша за решту.** {slow['Name']}: p95 = {ms(num(slow['95%']))} мс проти "
            f"{ms(num(light['95%']))} мс у GET /rooms/. Ймовірна причина (гіпотеза): хешування пароля bcrypt "
            "навмисно дороге й виконується синхронно всередині async-обробника, тож блокує event loop — "
            "під час хешування сервер не обслуговує інші запити."
        )
        advice.append(
            "Винести хешування/перевірку пароля з event loop (наприклад, `run_in_threadpool`), щоб важкі операції не блокували решту запитів."
        )

    if (
        light
        and (num(light["50%"]) or 0) > 0
        and (num(light["95%"]) or 0) > 5 * (num(light["50%"]) or 1)
    ):
        findings.append(
            f"**Довгий «хвіст» затримок навіть у найпростішого запиту.** GET /rooms/: медіана {ms(num(light['50%']))} мс, "
            f"p95 {ms(num(light['95%']))} мс, максимум {ms(num(light['Max Response Time']))} мс. Запити стоять у черзі за "
            "довшими операціями в одному процесі."
        )

    if app_cpu_max >= 90:
        findings.append(
            f"**Застосунок упирається в CPU** (максимум {app_cpu_max:.0f} % одного ядра"
            + (f", тоді як база даних — до {db_cpu_max:.1f} %" if by_name and db_cpu_max else "")
            + "). Один процес uvicorn виконує Python-код на одному ядрі, тому далі продуктивність зростає "
            "лише горизонтальним масштабуванням."
        )
        advice.append(
            "Запускати кілька воркерів (`uvicorn --workers N`) або кілька інстансів за балансувальником (як на схемі архітектури)."
        )
    elif by_name and db_cpu_max >= 90:
        findings.append(
            f"**База даних близька до насичення CPU** (до {db_cpu_max:.0f} %): потрібні індекси/оптимізація запитів або репліка для читання."
        )
        advice.append(
            "Додати індекси під запити перевірки слотів, винести читання в read-репліку, кешувати список залів у Redis."
        )

    if failures:
        top_err = sorted(failures, key=lambda r: -int(r["Occurrences"]))[:3]
        findings.append(
            "**Помилки під навантаженням:** "
            + "; ".join(f"{r['Name']} — {r['Error']} ({r['Occurrences']} разів)" for r in top_err)
            + "."
        )
        advice.append(
            "Розібрати причини помилок у логах застосунку (`docker compose logs app`) за часом їх появи на графіку."
        )
    else:
        findings.append("**Помилок (4xx/5xx, крім очікуваних 409) не зафіксовано.**")

    pool = server_log.get("pool_exhausted", 0)
    if pool:
        findings.append(
            f"**Вичерпання пулу з'єднань з БД.** У логах сервера {pool} помилок `QueuePool limit … connection timed out`. "
            "За замовчуванням SQLAlchemy тримає 5 з'єднань (+10 тимчасових) і чекає на вільне 30 с — звідси стеля "
            "затримки близько 30 с та відповіді 500. Ймовірна причина (гіпотеза): сесія БД відкрита на весь запит, "
            "а повільні операції (хешування пароля) тримають з'єднання, тож черга за з'єднаннями росте швидше, ніж вони звільняються."
        )
        advice.append(
            "Налаштувати пул (`pool_size`, `max_overflow`, `pool_timeout`) з урахуванням `max_connections` PostgreSQL і кількості воркерів; скоротити час утримання з'єднання."
        )
    elif server_log.get("unhandled_errors"):
        findings.append(
            f"**Необроблені помилки сервера (500):** {server_log['unhandled_errors']} — причини дивіться в логах застосунку за кодом `error_id`."
        )

    if double is not None:
        if double["overlapping_pairs"] > 0:
            ex = double["examples"][0] if double["examples"] else None
            findings.append(
                f"**Порушення цілісності даних (подвійне бронювання).** Після тесту серед {double['active_bookings']} активних "
                f"бронювань знайдено {double['overlapping_pairs']} пар із перетином часу в {double['rooms_affected']} залах"
                + (f" (приклад: «{ex['room']}», {ex['start']})" if ex else "")
                + ". Перевірка перетину й запис не атомарні: два одночасні запити обидва проходять перевірку (race condition). "
                "Це найсерйозніша знахідка для предметної області."
            )
            advice.append(
                "Зробити перевірку слота й запис атомарними: блокування в Redis (SET NX з TTL) або advisory lock / exclusion constraint у PostgreSQL."
            )
        else:
            findings.append(
                f"**Подвійних бронювань не виявлено** серед {double['active_bookings']} активних. "
                "Це не доказ відсутності гонки: при більшому навантаженні на один слот її варто перевірити повторно."
            )

    rep.h(2, "7. Виявлені вузькі місця")
    rep.bullets(findings)

    rep.h(2, "8. Висновки та подальші кроки")
    rep.p(
        f"За профілем «{m['profile_title']}» система обробила {requests_total} запитів із середнім "
        f"{requests_total / duration:.1f} RPS (пік {peak_rps:.1f}); 95-й перцентиль часу відповіді — "
        f"{ms(num(total['95%']))} мс, error rate — {error_rate:.2f} %."
    )
    advice.append(
        "Повторити тест після оптимізацій з тим самим профілем і порівняти метрики «до/після»."
    )
    rep.bullets(advice)

    (res / "report.md").write_text(rep.to_md(), encoding="utf-8")
    rep.to_docx(res / "report.docx")
    print(f"[report] готово: {res / 'report.docx'} та {res / 'report.md'}")


if __name__ == "__main__":
    main()
